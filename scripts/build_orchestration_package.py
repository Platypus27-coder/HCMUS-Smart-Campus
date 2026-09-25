"""Export authored subspace 06 sources to Word/Excel; no runtime integration.

Install requirements-orchestration.txt before running. Original documents are
archived separately and are never read or overwritten by this exporter.
"""
from pathlib import Path
import json
import re
import zipfile

import yaml
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'subspaces/06-central-orchestration'
IDS = ('campus-router-agent', 'campus-aggregator-agent')


def stable_office(path):
    """Normalize ZIP timestamps for reproducible builds from identical sources."""
    with zipfile.ZipFile(path) as source:
        entries = [(i.filename, source.read(i.filename)) for i in source.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            if name == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)', rb'\g<1>2000-01-01T00:00:00Z\2', data)
            item = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(item, data)


def export_doc(text, path):
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = 'HCMUS SMART CAMPUS  |  CENTRAL ORCHESTRATION'
    section.footer.paragraphs[0].text = 'Student-built prototype • Nội dung cấu hình và hướng dẫn'
    section.top_margin = section.bottom_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[name].font.color.rgb = RGBColor.from_string('146C70')
    code = False
    table = None
    for line in text.splitlines():
        if line.startswith('```'):
            code = not code
            table = None
            continue
        if not line.strip():
            continue
        if not code and line.startswith('|'):
            cells = [s.strip().replace('`', '') for s in line.strip('|').split('|')]
            if all(re.fullmatch(r':?-+:?', s) for s in cells):
                continue
            if table is None:
                table = doc.add_table(rows=0, cols=len(cells))
                table.style = 'Light Shading Accent 1'
            for cell, value in zip(table.add_row().cells, cells):
                cell.text = value
            continue
        table = None
        if code:
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
        elif line.startswith('# '):
            doc.add_heading(line[2:], 0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], 1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], 2)
        else:
            clean = line.replace('**', '').replace('`', '')
            doc.add_paragraph(clean[2:] if clean.startswith('- ') else clean,
                              style='List Bullet' if clean.startswith('- ') else None)
    doc.save(path)
    stable_office(path)


def build():
    directory = yaml.safe_load((ROOT / 'shared/agent-directory.yaml').read_text(encoding='utf-8'))['agents']
    graph = yaml.safe_load((ROOT / 'shared/agent-relationships.yaml').read_text(encoding='utf-8'))['relationships']
    for agent_id in IDS:
        folder = BASE / agent_id
        export_doc((folder / 'info.md').read_text(encoding='utf-8'), folder / 'info.docx')
        knowledge = (folder / 'agent/knowledge.md').read_text(encoding='utf-8')
        knowledge += '\n\n# Phụ lục — Danh mục chuẩn (sinh từ metadata)\n'
        for a in directory:
            knowledge += f"\n## {a['name']}\nID: {a['id']}\nSubspace: {a['subspace']}\nVai trò: {a['role']}\n"
            for key in ('capabilities', 'accepts', 'returns', 'out_of_scope', 'limitations'):
                knowledge += key + ': ' + ', '.join(a[key]) + '\n'
        knowledge += '\n# Phụ lục — Quan hệ bàn giao của agent\n'
        knowledge += '\n```yaml\n' + yaml.safe_dump(graph[agent_id], allow_unicode=True, sort_keys=False) + '```\n'
        for name in ('handoff-protocol.md', 'user-memory-policy.md'):
            knowledge += '\n# Phụ lục — ' + name + '\n' + (ROOT / 'shared' / name).read_text(encoding='utf-8')
        if agent_id == 'campus-router-agent':
            for name in ('routing-rules.yaml', 'routing-score-criteria.yaml'):
                knowledge += '\n# Phụ lục — ' + name + '\n```yaml\n' + (ROOT / 'shared' / name).read_text(encoding='utf-8') + '\n```\n'
        (folder / 'agent/knowledge-upload.md').write_text(knowledge, encoding='utf-8')
        export_doc(knowledge, folder / 'agent/knowledge.docx')
        rows = json.loads((folder / 'agent/qa.json').read_text(encoding='utf-8'))
        wb = Workbook()
        ws = wb.active
        ws.title = 'QA'
        ws.append(['Question', 'Answer', 'Display as Suggested Question?\nyes/no'])
        for row in rows:
            ws.append([row['question'], row['answer'], 'yes' if row.get('suggested', False) else 'no'])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column, width in [('A', 58), ('B', 110), ('C', 24)]:
            ws.column_dimensions[column].width = width
        for row in ws:
            ws.row_dimensions[row[0].row].height = 75 if row[0].row > 1 else 34
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == 1:
                    cell.font = Font(color='FFFFFF', bold=True)
                    cell.fill = PatternFill('solid', fgColor='146C70')
        wb.save(folder / 'agent/QA.xlsx')
        stable_office(folder / 'agent/QA.xlsx')
        print(f'[OK] {agent_id}: {len(rows)} Q&A rows, Word exports')
    for path in BASE.glob('*/menu/file/*.md'):
        export_doc(path.read_text(encoding='utf-8'), path.with_suffix('.docx'))
    # Explicit allowlist keeps archives and user-created numbered versions out.
    bundle = BASE / 'deployment/subspace-06-upload.zip'
    with zipfile.ZipFile(bundle, 'w', zipfile.ZIP_DEFLATED) as target:
        for aid in IDS:
            folder = BASE / aid
            config = json.loads((folder/'space/space_config.json').read_text(encoding='utf-8'))
            paths = ['info.docx', 'system-prompt.md', 'agent/knowledge.docx',
                     'agent/QA.xlsx', 'agent/background.png', 'space/welcome.md',
                     'space/space_config.json']
            paths += config['menu_resources']['files'] + config['menu_resources']['guides']
            for name in paths:
                target.write(folder/name, f'{aid}/{name}')
        for name in ('deployment/UPLOAD.md', 'deployment/agent-links.json',
                     'evaluation/demo-script.md', 'evaluation/partial-responses.example.json',
                     'evaluation/cases.json', 'evaluation/results-template.csv'):
            target.write(BASE/name, name)
    stable_office(bundle)
    print('[OK] Upload handover ZIP (extract and configure manually; not a native import format)')


if __name__ == '__main__':
    build()
